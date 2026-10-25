"""Isolated extraction: a parser hang or crash fails one source, not the queue.

`ingest._extract_isolated` runs `extract_sections` in a child process
(`app.extract_child`). These tests drive the real child for the round-trip
contract, and replace the child's command line (`_extract_child_argv`) only to
simulate a hang or a crash, which no real fixture can do on demand.
"""
import asyncio
import io
import json
import os
import sys
import time
from pathlib import Path

import pytest


# -------------------- fixtures --------------------

def _minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def _write_fixture(kind: str, directory: Path) -> Path:
    if kind == "pdf":
        path = directory / "report.pdf"
        path.write_bytes(_minimal_pdf("Quarterly revenue grew 12 percent"))
    elif kind == "docx":
        import docx
        path = directory / "memo.docx"
        document = docx.Document()
        document.add_heading("Project Alpha", level=1)
        document.add_paragraph("Alpha revenue is 42 dollars.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "region", "sales"
        table.cell(1, 0).text, table.cell(1, 1).text = "north", "7"
        document.save(path)
    elif kind == "pptx":
        from pptx import Presentation
        path = directory / "deck.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Roadmap"
        slide.placeholders[1].text = "Ship isolation first"
        slide.notes_slide.notes_text_frame.text = "speaker note"
        deck.save(path)
    elif kind == "xlsx":
        import openpyxl
        path = directory / "sales.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Q3"
        sheet.append(["region", "sales"])
        sheet.append(["north", 7])
        sheet.append(["south", 9])
        book.save(path)
    elif kind == "csv":
        path = directory / "sales.csv"
        # Enough rows for encoding detection to settle; a 3-line Big5 file is
        # misdetected today (separate bug), which would make this vacuous.
        rows = "".join(f"北區分公司{i},{i},王小明\n" for i in range(40))
        path.write_bytes(("地區,銷售,負責人\n" + rows).encode("cp950"))
    elif kind == "vtt":
        path = directory / "meeting.vtt"
        path.write_text(
            "WEBVTT\n\n00:00:01.000 --> 00:00:04.000\n<v 王小明>我們下週上線。\n",
            encoding="utf-8",
        )
    elif kind == "html":
        path = directory / "page.html"
        path.write_text("<html><body><h1>Title</h1><p>Body text here.</p></body></html>", encoding="utf-8")
    else:
        path = directory / "notes.txt"
        path.write_text("Alpha project revenue is 42 dollars.", encoding="utf-8")
    return path


def _insert_source(db, path: Path) -> int:
    with db.connect() as conn:
        user = conn.execute("SELECT * FROM users WHERE username = 'user'").fetchone()
        return conn.execute(
            """
            INSERT INTO sources (user_id, filename, stored_path, content_type, status)
            VALUES (?, ?, ?, 'application/octet-stream', 'uploaded')
            """,
            (user["id"], path.name, str(path)),
        ).lastrowid


def _source(db, source_id: int):
    with db.connect() as conn:
        return conn.execute("SELECT * FROM sources WHERE id = ?", (source_id,)).fetchone()


def _hang_argv(pid_file: Path) -> list[str]:
    """A child that records its pid, then spins forever (a hung parser)."""
    script = (
        "import os, sys\n"
        f"open({str(pid_file)!r}, 'w').write(str(os.getpid()))\n"
        "while True:\n    pass\n"
    )
    return [sys.executable, "-c", script]


def _pid_alive(pid: int) -> bool:
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    return True


def _wait_for(predicate, timeout: float = 20.0) -> bool:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if predicate():
            return True
        time.sleep(0.05)
    return False


# -------------------- G5 / G10: the producer -> reader contract --------------------

@pytest.mark.parametrize("kind", ["pdf", "docx", "pptx", "xlsx", "csv", "vtt", "html", "txt"])
def test_isolated_result_equals_in_process_result(fresh_modules, tmp_path, kind):
    """The real child must return exactly what in-process extraction returns."""
    ingest = fresh_modules.ingest
    path = _write_fixture(kind, tmp_path)

    expected = ingest.extract_sections(path)
    isolated = asyncio.run(ingest._extract_isolated(path))

    assert expected.sections, f"fixture for {kind} extracted nothing; the comparison would be vacuous"
    assert isolated == expected


def test_extractor_error_message_survives_the_process_boundary(fresh_modules, tmp_path):
    """`sources.error` shows `str(exc)`, so the child must not reword it."""
    ingest = fresh_modules.ingest
    path = tmp_path / "broken.docx"
    path.write_bytes(b"this is not a zip archive")

    with pytest.raises(Exception) as in_process:
        ingest.extract_sections(path)
    with pytest.raises(ingest.ExtractionError) as isolated:
        asyncio.run(ingest._extract_isolated(path))

    assert str(isolated.value) == str(in_process.value)


def test_child_starts_when_the_parent_runs_from_another_directory(fresh_modules, tmp_path, monkeypatch):
    """`python -m app.extract_child` must resolve even if cwd is not the app root
    (e.g. `uvicorn --app-dir`), without the operator setting PYTHONPATH."""
    ingest = fresh_modules.ingest
    path = _write_fixture("txt", tmp_path)
    elsewhere = tmp_path / "elsewhere"
    elsewhere.mkdir()
    monkeypatch.chdir(elsewhere)
    monkeypatch.delenv("PYTHONPATH", raising=False)

    assert asyncio.run(ingest._extract_isolated(path)).sections


def test_child_reads_the_same_config_as_the_parent(fresh_modules, tmp_path, monkeypatch):
    """Extract-time caps come from env/config.toml, which the child resolves itself."""
    ingest = fresh_modules.ingest
    path = _write_fixture("xlsx", tmp_path)
    monkeypatch.setenv("NOTEBOOKLM_RUNTIME_EXTRACT_MAX_FILE_BYTES", "10")

    with pytest.raises(ingest.ExtractionError, match="too large to ingest"):
        asyncio.run(ingest._extract_isolated(path))


# -------------------- G1 / G3 / G4 / G8: hang, kill, crash --------------------

def test_hung_parser_fails_the_source_and_the_child_is_killed(fresh_modules, local_embed, tmp_path, monkeypatch):
    db, ingest = fresh_modules.db, fresh_modules.ingest
    pid_file = tmp_path / "child.pid"
    monkeypatch.setattr(ingest.config.jobs, "extract_timeout_s", 1.0)
    monkeypatch.setattr(ingest, "_extract_child_argv", lambda *a: _hang_argv(pid_file))
    source_id = _insert_source(db, _write_fixture("txt", tmp_path))

    started = time.monotonic()
    asyncio.run(ingest.process_source(source_id))
    elapsed = time.monotonic() - started

    row = _source(db, source_id)
    diagnostics = json.loads(row["diagnostics_json"])
    assert row["status"] == "failed"
    assert diagnostics["failed_stage"] == "extract"
    assert diagnostics["failed_reason"] == ingest.EXTRACT_FAILED_TIMEOUT
    assert "1 秒" in row["error"]
    assert elapsed < 6, f"timeout of 1s took {elapsed:.1f}s to fail the source"
    assert not _pid_alive(int(pid_file.read_text())), "hung child outlived its timeout"


def test_cancelling_extraction_kills_the_child(fresh_modules, tmp_path, monkeypatch):
    """Worker shutdown cancels the task mid-parse; the child must die with it."""
    ingest = fresh_modules.ingest
    pid_file = tmp_path / "child.pid"
    monkeypatch.setattr(ingest.config.jobs, "extract_timeout_s", 60.0)
    monkeypatch.setattr(ingest, "_extract_child_argv", lambda *a: _hang_argv(pid_file))

    async def drive():
        task = asyncio.create_task(ingest._extract_isolated(tmp_path / "any.txt"))
        while not pid_file.exists() or not pid_file.read_text():
            await asyncio.sleep(0.05)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task

    asyncio.run(asyncio.wait_for(drive(), timeout=20))
    assert not _pid_alive(int(pid_file.read_text())), "cancelled extraction left its child running"


def test_crashing_child_fails_the_source_with_crash_reason(fresh_modules, local_embed, tmp_path, monkeypatch):
    db, ingest = fresh_modules.db, fresh_modules.ingest
    monkeypatch.setattr(
        ingest, "_extract_child_argv", lambda *a: [sys.executable, "-c", "import os; os.abort()"]
    )
    source_id = _insert_source(db, _write_fixture("txt", tmp_path))

    asyncio.run(ingest.process_source(source_id))

    row = _source(db, source_id)
    assert row["status"] == "failed"
    assert json.loads(row["diagnostics_json"])["failed_reason"] == ingest.EXTRACT_FAILED_CRASHED


def test_queue_keeps_draining_behind_a_hung_source(fresh_modules, local_embed, tmp_path, monkeypatch):
    """The user-visible property: one hung upload must not block everyone else's."""
    import app.jobs as jobs
    import app.worker as worker

    db, ingest = fresh_modules.db, fresh_modules.ingest
    pid_file = tmp_path / "child.pid"
    real_argv = ingest._extract_child_argv
    monkeypatch.setattr(ingest.config.jobs, "extract_timeout_s", 1.0)
    monkeypatch.setattr(
        ingest,
        "_extract_child_argv",
        lambda path, *rest: _hang_argv(pid_file) if path.name == "hang.txt" else real_argv(path, *rest),
    )
    hung_path = tmp_path / "hang.txt"
    hung_path.write_text("never parsed", encoding="utf-8")
    hung_id = _insert_source(db, hung_path)
    good_id = _insert_source(db, _write_fixture("txt", tmp_path))
    jobs.enqueue_source(hung_id)
    jobs.enqueue_source(good_id)

    async def drive():
        stop = asyncio.Event()
        loop_task = asyncio.create_task(worker.run_worker_loop(poll_interval=0.05, stop_event=stop))
        deadline = time.monotonic() + 30
        while _source(db, good_id)["status"] != "indexed" and time.monotonic() < deadline:
            await asyncio.sleep(0.1)
        stop.set()
        await asyncio.wait_for(loop_task, timeout=10)

    asyncio.run(drive())

    assert _source(db, good_id)["status"] == "indexed", "the queue stalled behind the hung source"
    hung = _source(db, hung_id)
    assert hung["status"] == "failed"
    assert json.loads(hung["diagnostics_json"])["failed_reason"] == ingest.EXTRACT_FAILED_TIMEOUT
    with db.connect() as conn:
        attempts = conn.execute(
            "SELECT attempts FROM ingest_jobs WHERE source_id = ?", (hung_id,)
        ).fetchone()["attempts"]
    assert attempts == 1, "a timed-out source must not be retried"


# -------------------- G6: kill-switch, timeout clamp --------------------

def test_isolation_off_parses_in_process(fresh_modules, local_embed, tmp_path, monkeypatch):
    db, ingest = fresh_modules.db, fresh_modules.ingest
    monkeypatch.setattr(ingest.config.jobs, "extract_isolation", False)

    async def no_subprocess(*args, **kwargs):
        raise AssertionError("extract_isolation=false must not spawn a child")

    monkeypatch.setattr(ingest.asyncio, "create_subprocess_exec", no_subprocess)
    source_id = _insert_source(db, _write_fixture("txt", tmp_path))

    asyncio.run(ingest.process_source(source_id))

    assert _source(db, source_id)["status"] == "indexed"


def test_timeout_is_kept_below_the_job_visibility_timeout(fresh_modules, monkeypatch):
    ingest = fresh_modules.ingest
    monkeypatch.setattr(ingest.config.jobs, "visibility_timeout_s", 1800.0)

    monkeypatch.setattr(ingest.config.jobs, "extract_timeout_s", 300.0)
    assert ingest._extract_timeout_s() == 300.0

    monkeypatch.setattr(ingest.config.jobs, "extract_timeout_s", 5000.0)
    assert ingest._extract_timeout_s() == 1620.0
